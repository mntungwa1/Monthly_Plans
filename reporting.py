"""Validated draft model and template-based RMO Word exports."""
from copy import deepcopy
from datetime import date
from io import BytesIO
from pathlib import Path
import json
import math
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).parent
BASES = ['Actual', 'Estimated', 'Scheduled', 'Anticipated']
STATUSES = ['Planned', 'In progress', 'Completed', 'Delayed', 'Deferred']
FIELDS = ['phase', 'kra', 'activity', 'start', 'end', 'hours', 'basis', 'status',
          'output', 'purpose', 'alignment', 'evidence', 'owner', 'dependencies', 'notes']
NARRATIVES = ['summary', 'objectives', 'targets', 'outputs', 'integration', 'lessons',
              'conclusion', 'annexures', 'plan_summary', 'plan_objectives', 'plan_targets',
              'plan_integration', 'plan_risks', 'plan_conclusion']

def new_draft():
    today = date.today()
    first = today.replace(day=1)
    import calendar
    last = today.replace(day=calendar.monthrange(today.year, today.month)[1])
    return {'schema_version': 1, 'meta': {'substream': '', 'prepared_by': '',
        'role': '', 'report_start': first.isoformat(), 'report_end': today.isoformat(),
        'plan_start': today.isoformat(), 'plan_end': last.isoformat(),
        'signoff_date': today.isoformat(), 'monthly_limit': 96.0,
        'hours_note': '', 'contract_note': '', 'logo_variant': 'Enhanced'},
        'narratives': {k: '' for k in NARRATIVES}, 'activities': []}

def clean_text(value):
    if value is None: return ''
    if isinstance(value, float) and math.isnan(value): return ''
    if hasattr(value, 'isoformat'): return value.isoformat()
    return str(value).strip()

def normalise_rows(records):
    out = []
    for row in records:
        if not any(clean_text(row.get(k)) for k in ['activity', 'output', 'kra']):
            continue
        r = {k: clean_text(row.get(k)) for k in FIELDS}
        raw = row.get('hours', 0)
        r['hours'] = 0.0 if raw is None or clean_text(raw) == '' else float(raw)
        r['phase'] = r['phase'] or 'Progress'
        r['basis'] = r['basis'] or ('Anticipated' if r['phase'] == 'Workplan' else 'Estimated')
        r['status'] = r['status'] or 'Planned'
        out.append(r)
    return out

def validate(draft):
    errors, warnings = [], []
    m = draft['meta']
    for k in ['substream', 'prepared_by']:
        if not m[k].strip(): errors.append(f"Enter {k.replace('_', ' ')}.")
    for prefix in ['report', 'plan']:
        if date.fromisoformat(m[prefix+'_start']) > date.fromisoformat(m[prefix+'_end']):
            errors.append(f'{prefix.title()} start must not be after its end.')
    if not draft['activities']: warnings.append('No activities have been captured.')
    for i, r in enumerate(draft['activities'], 1):
        label = f'Activity {i}'
        if r['phase'] not in ['Progress', 'Workplan']: errors.append(f'{label}: invalid phase.')
        if r['basis'] not in BASES: errors.append(f'{label}: invalid hours basis.')
        if r['status'] not in STATUSES: errors.append(f'{label}: invalid status.')
        if not math.isfinite(r['hours']) or r['hours'] < 0: errors.append(f'{label}: hours must be finite and non-negative.')
        if not r['activity']: errors.append(f'{label}: enter an activity name.')
        if not r['start'] or not r['end']:
            errors.append(f'{label}: enter start and end dates.'); continue
        try:
            a,b = date.fromisoformat(r['start']),date.fromisoformat(r['end'])
        except ValueError:
            errors.append(f'{label}: use valid dates.'); continue
        if a > b: errors.append(f'{label}: end date precedes start date.')
        prefix = 'plan' if r['phase'] == 'Workplan' else 'report'
        if a < date.fromisoformat(m[prefix+'_start']) or b > date.fromisoformat(m[prefix+'_end']):
            warnings.append(f'{label}: dates fall outside the {prefix} window; review before export.')
        if r['phase'] == 'Workplan' and r['basis'] == 'Actual':
            errors.append(f'{label}: actual hours belong in Progress; workplan hours are planned.')
        if r['basis'] == 'Actual' and not r['evidence']:
            errors.append(f'{label}: actual hours require a supporting evidence reference.')
        if r['status'] == 'Completed' and not r['evidence']:
            warnings.append(f'{label}: completed status has no evidence reference.')
        if not r['kra'] or not r['output']:
            warnings.append(f'{label}: add the KRA and deliverable.')
    plan_hours = sum(r['hours'] for r in draft['activities'] if r['phase']=='Workplan')
    if plan_hours > m['monthly_limit'] and m['monthly_limit'] > 0:
        warnings.append(f'Workplan hours ({plan_hours:g}) exceed the entered monthly ceiling ({m["monthly_limit"]:g}).')
    return errors, warnings

def loads_draft(raw):
    if len(raw) > 2_000_000: raise ValueError('Draft is too large (maximum 2 MB).')
    def reject_constant(value):
        raise ValueError(f'Invalid JSON number: {value}')
    d = json.loads(raw,parse_constant=reject_constant)
    if not isinstance(d, dict) or d.get('schema_version') != 1:
        raise ValueError('Use a version 1 RMO app draft JSON file.')
    base = new_draft()
    if not isinstance(d.get('meta'),dict) or not isinstance(d.get('narratives'),dict) or not isinstance(d.get('activities'),list):
        raise ValueError('Draft must contain metadata, narratives and an activity list.')
    if len(d['activities']) > 500: raise ValueError('Maximum 500 activities per draft.')
    for k in base['meta']:
        if k in d['meta']: base['meta'][k] = clean_text(d['meta'][k])
    base['meta']['monthly_limit'] = float(base['meta']['monthly_limit'])
    if not math.isfinite(base['meta']['monthly_limit']) or base['meta']['monthly_limit'] < 0:
        raise ValueError('Monthly ceiling must be finite and non-negative.')
    for k in ['report_start','report_end','plan_start','plan_end','signoff_date']:
        date.fromisoformat(base['meta'][k])
    for k in NARRATIVES: base['narratives'][k] = clean_text(d['narratives'].get(k))
    if any(not isinstance(r,dict) for r in d['activities']): raise ValueError('Each activity must be a record.')
    base['activities'] = normalise_rows(d['activities'])
    for r in base['activities']:
        if r['phase'] not in ['Progress','Workplan'] or r['basis'] not in BASES or r['status'] not in STATUSES:
            raise ValueError('Draft contains an invalid phase, hours basis or status.')
        if not math.isfinite(r['hours']) or r['hours'] < 0: raise ValueError('Invalid activity hours.')
        for k in ['start','end']:
            if r[k]: date.fromisoformat(r[k])
    return base

def period(m, prefix):
    return f"{date.fromisoformat(m[prefix+'_start']):%d %B %Y} to {date.fromisoformat(m[prefix+'_end']):%d %B %Y}"

def totals(rows):
    return {b:sum(r['hours'] for r in rows if r['basis']==b) for b in BASES}

def hour_line(rows):
    t = totals(rows)
    return '; '.join(f'{b}: {t[b]:g} hours' for b in BASES) + '. Actual hours exclude estimates, scheduled durations and anticipated work.'

def sections(draft, mode):
    m,n = draft['meta'],draft['narratives']
    planned = mode == 'workplan'
    phase = 'Workplan' if planned else 'Progress'
    rows = [r for r in draft['activities'] if r['phase']==phase]
    key = lambda x: n.get('plan_'+x if planned else x,'')
    empty = 'No narrative entered.'
    output = '\n'.join(f"{i}. {r['output'] or r['activity']} — {r['status']}. {r['evidence'] or 'Evidence reference not entered.'}" for i,r in enumerate(rows,1))
    return rows, [
        ('1. Executive Summary', key('summary') or empty),
        ('2. Objectives for the Period', key('objectives') or empty),
        ('3. Summary of Performance Targets', (key('targets') or empty)+'\n'+hour_line(rows)+'\n'+(m['hours_note'] or '')+'\n'+(m['contract_note'] or '')),
        ('4. Summary of Expected Outputs' if planned else '4. Summary of Key Outputs', (output or 'No outputs captured.') if planned else (n['outputs'] or output or 'No outputs captured.')),
        ('5. Detailed Planned Activities' if planned else '5. Detailed Activities Undertaken', ''),
        ('6. Cross Stream Integration', key('integration') or empty),
        ('7. Dependencies and Delivery Risks' if planned else '7. Key Lessons Emerging Issues and Risks', key('risks') if planned else n['lessons']),
        ('8. Deliverables Against the Workplan' if planned else '8. Outputs Against the Workplan', ''),
        ('9. Conclusion', key('conclusion') or empty),
        ('10. Workplan Review and Sign Off' if planned else '10. Report Affirmation and Sign Off',
         f"This workplan records intended activities for the {m['substream']} sub-stream for {period(m,'plan')}. Hours are planning allocations subject to review and approval." if planned else
         f"I confirm, upon verification and signing below, that this report records my contributions to the {m['substream']} sub-stream for {period(m,'report')}. Actual, estimated, scheduled and anticipated hours are identified separately. Evidence references and completion claims must be verified before submission.")]

def export_docx(draft, mode='progress'):
    errors,_ = validate(draft)
    if errors: raise ValueError('\n'.join(errors))
    d = Document(ROOT/'templates/rmo_template.docx')
    originals = [deepcopy(p._p) for p in d.paragraphs]
    tbl = deepcopy(d.tables[0]._tbl)
    # Keep the original section geometry, styles, headers, footers and media relationships.
    body = d._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'): body.remove(child)
    for s in ['Heading 1','Heading 2']:
        d.styles[s].font.color.rgb=RGBColor(0,0,0)
    def insert(el):
        sect = body.find(qn('w:sectPr'))
        if sect is None: body.append(el)
        else: sect.addprevious(el)
    def para(text, kind=18):
        el=deepcopy(originals[kind]); ppr=el.find(qn('w:pPr'))
        rpr=None
        first=el.find(qn('w:r'))
        if first is not None and first.find(qn('w:rPr')) is not None: rpr=deepcopy(first.find(qn('w:rPr')))
        for child in list(el):
            if child is not ppr: el.remove(child)
        r=OxmlElement('w:r')
        if rpr is not None:
            for c in list(rpr):
                if c.tag == qn('w:color'): rpr.remove(c)
            r.append(rpr)
        for j,line in enumerate(text.split('\n')):
            if j: r.append(OxmlElement('w:br'))
            t=OxmlElement('w:t');t.set(qn('xml:space'),'preserve');t.text=line;r.append(t)
        el.append(r);insert(el)
        from docx.text.paragraph import Paragraph
        p=Paragraph(el,d._body)
        p.paragraph_format.space_after=Pt(6)
        p.paragraph_format.line_spacing=1.08
        if '\n' in text and kind not in [2,3,4,5,11]:
            p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        if kind in [17,32]:
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_before=Pt(10)
        else:p.paragraph_format.keep_with_next=False
        return p
    # Prominent institutional branding, with original asset retained as fallback.
    logo=ROOT/'assets/logos_enhanced.png' if draft['meta'].get('logo_variant')=='Enhanced' else ROOT/'assets/image1.png'
    if not logo.exists(): logo=ROOT/'assets/image1.png'
    p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(logo),width=Inches(5.8))
    para('NATIONAL DISASTER MANAGEMENT CENTRE\nDepartment of Cooperative Governance',2)
    para('MONTHLY WORKPLAN' if mode=='workplan' else 'EXECUTIVE SUMMARY OF PORTFOLIO OF EVIDENCE',3)
    m=draft['meta']
    para('Sub-stream: '+m['substream'],4)
    para(('Workplan Period: '+period(m,'plan')) if mode=='workplan' else ('Reporting Period: '+period(m,'report')),5)
    para('Prepared by: '+m['prepared_by']+ ('\n'+m['role'] if m['role'] else '')+'\nDisaster Management Results Management Office Stream',11)
    rows,content=sections(draft,mode)
    for heading,text in content:
        para(heading,17)
        if heading.startswith('5.'):
            if not rows: para('No activities captured.')
            for i,r in enumerate(rows,1):
                para(f"5.{i} {r['activity']}",32)
                para(f"Dates: {r['start']} to {r['end']} | KRA: {r['kra'] or 'Not entered'}\nHours: {r['hours']:g} ({r['basis'].lower()}) | Status: {r['status']}\nOwner: {r['owner'] or m['prepared_by']}")
                for label,k in [('Output','output'),('Purpose','purpose'),('Strategy and APP alignment','alignment'),('Evidence','evidence'),('Dependencies','dependencies'),('Activity detail','notes')]:
                    if r[k]: para(f'{label}: {r[k]}')
        elif heading.startswith('8.'):
            from docx.table import Table
            insert(deepcopy(tbl));t=Table(body[-2],d._body)
            for rr in list(t._tbl.tr_lst)[1:]:t._tbl.remove(rr)
            headers=['Output','Purpose','Contribution to DM RMO Strategy and Matrix and NDMC APP Targets','Relevant PoE and status']
            for c,txt in zip(t.rows[0].cells,headers): c.text=txt
            for r in rows:
                cells=t.add_row().cells
                vals=[r['output'] or r['activity'],r['purpose'],f"{r['kra']}\n{r['alignment']}",f"{r['status']}\n{r['evidence'] or 'Evidence not entered'}\n{r['hours']:g} hours ({r['basis'].lower()})"]
                for c,v in zip(cells,vals):c.text=v
            trpr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');trpr.append(repeat)
            for ri,rr in enumerate(t.rows):
                trpr=rr._tr.get_or_add_trPr()
                if trpr.find(qn('w:cantSplit')) is None: trpr.append(OxmlElement('w:cantSplit'))
                for c in rr.cells:
                    for p in c.paragraphs:
                        p.paragraph_format.space_after=Pt(5);p.paragraph_format.keep_with_next=False
                        for run in p.runs:run.font.size=Pt(10);run.font.bold=(ri==0)
                for ht in rr._tr.xpath('./w:trPr/w:trHeight'):ht.getparent().remove(ht)
        else:
            for line in (text or 'No narrative entered.').split('\n'):
                if line.strip():para(line)
        if heading.startswith('10.'):
            para('Signature __________________\nName: '+m['prepared_by']+'\nDate: '+m['signoff_date'],83)
    para('ANNEXURES',17)
    for line in (draft['narratives']['annexures'] or 'No annexure references entered.').split('\n'):
        if line.strip():para(line)
    d.core_properties.title=f"{m['substream']} Monthly {'Workplan' if mode=='workplan' else 'Progress Report'}"
    d.core_properties.author=m['prepared_by']
    d.core_properties.comments='Generated using the supplied DM RMO reporting template.'
    out=BytesIO();d.save(out);return out.getvalue()

def filename(draft,mode):
    stem=re.sub(r'[^A-Za-z0-9_-]+','_',draft['meta']['substream']).strip('_') or 'RMO'
    key='plan_start' if mode=='workplan' else 'report_start'
    return f"{stem}_{mode}_{draft['meta'][key][:7]}.docx"
