import json
import sys
import unittest
from pathlib import Path
from io import BytesIO
from copy import deepcopy

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from reporting import *
from docx import Document

class ReportingTests(unittest.TestCase):
    def setUp(self):self.d=loads_draft((ROOT/'examples/bbb_example.json').read_bytes())
    def test_separate_hours(self):
        t=totals([r for r in self.d['activities'] if r['phase']=='Progress'])
        self.assertEqual(t,{'Actual':0,'Estimated':48,'Scheduled':7,'Anticipated':0})
        self.assertEqual(sum(r['hours'] for r in self.d['activities'] if r['phase']=='Workplan'),72)
    def test_draft_roundtrip(self):
        self.assertEqual(self.d,loads_draft(json.dumps(self.d).encode()))
    def test_actual_requires_evidence(self):
        r=self.d['activities'][0];r['basis']='Actual';r['evidence']=''
        self.assertTrue(any('evidence' in x for x in validate(self.d)[0]))
    def test_invalid_dates_and_hours(self):
        self.d['activities'][0]['hours']=-1
        self.d['activities'][1]['end']='2026-09-01'
        self.assertEqual(len(validate(self.d)[0]),2)
        with self.assertRaises(ValueError):export_docx(self.d)
    def test_json_invalid_numbers(self):
        self.d['activities'][0]['hours']=float('nan')
        with self.assertRaises(ValueError):loads_draft(json.dumps(self.d).encode())
    def test_workplan_blocks_actual(self):
        self.d['activities'][-1]['basis']='Actual'
        self.d['activities'][-1]['evidence']='Some record'
        self.assertTrue(any('belong in Progress' in x for x in validate(self.d)[0]))
    def test_word_exports(self):
        for mode,count in [('progress',6),('workplan',4)]:
            doc=Document(BytesIO(export_docx(self.d,mode)))
            self.assertEqual(len(doc.tables[0].rows),count+1)
            self.assertEqual(len(doc.tables[0].columns),4)
            self.assertEqual(len(doc.inline_shapes),1)
            self.assertIn('Confidential',doc.sections[0].footer.paragraphs[0].text)
            text='\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('10.',text)
            self.assertNotIn('……………………',text)
            if mode=='progress':self.assertNotIn('5.7',text)

class AppTests(unittest.TestCase):
    def test_blank_example_and_restored_widgets(self):
        from streamlit.testing.v1 import AppTest
        at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30).run()
        self.assertFalse(at.exception)
        next(b for b in at.button if b.label=='Load BBB example').click().run()
        self.assertFalse(at.exception)
        self.assertEqual(at.text_input[0].value,'Build Back Better (BBB)')
        self.assertEqual([x.value for x in at.metric],['0 h','48 h','7 h','0 h','72 h'])
        # Restore the same way the upload button does: replace draft and increment widget epoch.
        restored=loads_draft(json.dumps(at.session_state['draft']).encode())
        restored['meta']['substream']='SEMS'
        restored['narratives']['summary']='Restored summary'
        at.session_state['draft']=restored
        at.session_state['epoch']+=1
        del at.session_state['editor_source']
        at.run()
        self.assertFalse(at.exception)
        self.assertEqual(at.text_input[0].value,'SEMS')
        self.assertEqual(at.text_area[2].value,'Restored summary')
    def test_password_gate(self):
        from streamlit.testing.v1 import AppTest
        at=AppTest.from_file(str(ROOT/'app.py'),default_timeout=30)
        at.secrets['APP_PASSWORD']='test-password'
        at.run();self.assertFalse(at.exception)
        self.assertEqual(at.text_input[0].label,'App password')
        at.text_input[0].set_value('test-password')
        next(b for b in at.button if b.label=='Sign in').click().run()
        self.assertFalse(at.exception)
        self.assertEqual(at.text_input[0].label,'Sub-stream name')

if __name__=='__main__':unittest.main()
